"""Verify cross-artifact consistency, then package only submission files."""
import csv
import ast
import hashlib
import json
import re
import zipfile
from pathlib import Path

import numpy as np
import openpyxl

from results_table import load_results


def main():
    root = Path(__file__).resolve().parents[1]
    repo = root.parent
    summary = json.loads((root / 'results/part4_summary.json').read_text())
    manifest = json.loads((root / 'results/evaluation_manifest.json').read_text())
    scores = json.loads((root / 'eval_result.json').read_text())
    results = load_results(str(root / 'results'))
    ids = {r['cfg']['exp_id'] for r in results}
    assert len(ids) == len(results) == summary['experiment_count']
    assert {'loss', 'optimizer', 'hparam', 'dropout', 'clipping', 'amp', 'init'} <= {r['cfg']['group'] for r in results}
    assert len(list(root.glob('**/*.pt'))) == 0
    for result in results:
        assert (root / f"figures/{result['cfg']['exp_id']}.png").is_file()
        assert result['provenance']['train_rows'] == 371847
        assert result['provenance']['val_rows'] == 92962
        assert result['summary']['diverged'] or len(result['history']['epoch']) == result['cfg']['epochs']
    for name, pred, score in [('baseline', root / 'results/predictions_baseline.csv', root / 'results/baseline_eval_result.json'),
                              ('final', root / 'predictions_eval.csv', root / 'eval_result.json')]:
        assert hashlib.sha256(pred.read_bytes()).hexdigest() == manifest[name]['pred_sha256']
        assert hashlib.sha256(score.read_bytes()).hexdigest() == manifest[name]['scores_sha256']
    with (root / 'predictions_eval.csv').open() as stream:
        predictions = list(csv.DictReader(stream))
    with np.load(repo / 'data/processed/eval.npz') as data:
        assert {int(row['row_id']) for row in predictions} == set(data['row_id'])
    assert len(predictions) == scores['n_eval'] == 116203
    assert all(0 <= int(row['pred']) < 7 for row in predictions)
    workbook = openpyxl.load_workbook(root / 'experiments.xlsx')
    assert workbook.sheetnames == ['Legend', 'Experiments', 'Seeds', 'Summary']
    sheet = workbook['Experiments']
    headers = [c.value for c in sheet[1]]
    table = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    assert {r['exp_id'] for r in table} == ids
    result_map = {r['cfg']['exp_id']: r for r in results}
    for row in table:
        expected_f1 = result_map[row['exp_id']]['summary']['val_macro_f1']
        assert (row['val_macro_f1'] is None and expected_f1 is None) or abs(row['val_macro_f1'] - expected_f1) < 1e-12
        assert row['delta_val_f1_vs_base'].startswith('=')
        if row['eval_macro_f1'] is not None:
            assert row['group'] in ('baseline', 'final')
    final_row = next(r for r in table if r['exp_id'] == summary['final_exp_id'])
    assert abs(final_row['eval_macro_f1'] - scores['macro_f1']) < 1e-12
    assert scores['macro_f1'] == summary['final_eval_macro_f1']
    assert abs(final_row['eval_acc'] - scores['accuracy']) < 1e-12
    cached = openpyxl.load_workbook(root / 'experiments.xlsx', data_only=True)
    formula_count = 0
    for formula_sheet in workbook:
        for line in formula_sheet:
            for cell in line:
                assert cached[formula_sheet.title][cell.coordinate].data_type != 'e'
                if cell.data_type == 'f':
                    formula_count += 1
    # Recalculation is required before packaging, not only formulas being present.
    assert abs(cached['Seeds']['C8'].value - np.mean([r['summary']['val_macro_f1'] for r in results if r['cfg']['group'] == 'baseline'])) < 1e-12
    report = (root / 'REPORT.md').read_text(encoding='utf-8')
    for target in re.findall(r'\]\(([^)]+)\)', report):
        assert (root / target).exists(), target
    notebook = json.loads((root / 'code/lab.ipynb').read_text(encoding='utf-8'))
    for cell in notebook['cells']:
        if cell['cell_type'] == 'code':
            assert cell['execution_count'] is not None
            assert not any(o['output_type'] == 'error' for o in cell['outputs'])
    for path in (root / 'code').glob('*.py'):
        for node in ast.walk(ast.parse(path.read_text(encoding='utf-8'))):
            if isinstance(node, ast.Raise):
                exception = node.exc.func if isinstance(node.exc, ast.Call) else node.exc
                assert not (isinstance(exception, ast.Name) and exception.id == 'NotImplementedError'), path
    archive = repo / (root.name + '.zip')
    with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as bundle:
        for path in sorted(root.rglob('*')):
            if not path.is_file() or '__pycache__' in path.parts or path.name == 'part0_colab_cuda.ipynb':
                continue
            bundle.write(path, Path(root.name) / path.relative_to(root))
    print(f'PASS: {len(results)} experiment rows and plots; 7 topics; 116203 predictions; {formula_count} formulas; notebook complete.')
    print('Package:', archive)


if __name__ == '__main__':
    main()
